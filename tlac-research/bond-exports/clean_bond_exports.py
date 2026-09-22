"""Clean Workspace bond exports locally; no LSEG requests are made."""

# ---- EDIT THIS --------------------------------------------------------------
input_folder = "bond_exports"
output_folder = "bond_exports_cleaned"
# -----------------------------------------------------------------------------
try:
    from config import DATA_DIR
except ImportError:
    raise RuntimeError(
        "config.py not found. Copy config.example.py to config.py and set DATA_DIR."
    ) from None

directory = DATA_DIR.replace("\\", "/")

import os
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import load_workbook


META = ["source_file", "source_excel_row", "export_isin", "count_matches_total"]
MONTHS = {name: i for i, name in enumerate(
    "jan feb mar apr may jun jul aug sep oct nov dec".split(), 1
)}


def populated(value):
    return value is not None and str(value).strip() != ""


def keep_maturity(value):
    """Accept typed Excel dates, English dd-Mon-yyyy text, and Perpetual."""
    if isinstance(value, (date, datetime)):
        return True
    if not isinstance(value, str):
        return False
    value = value.strip()
    if value.casefold() == "perpetual":
        return True
    match = re.fullmatch(r"(\d{2})-([A-Za-z]{3})-(\d{4})", value)
    if not match:
        return False
    day, month, year = match.groups()
    try:
        date(int(year), MONTHS[month.lower()], int(day))
        return True
    except (KeyError, ValueError):
        return False


def parse_total(value):
    try:
        count = Decimal(str(value).strip().replace(",", ""))
        if not count.is_finite() or count < 0 or count != count.to_integral_value():
            raise ValueError
        return int(count)
    except (InvalidOperation, ValueError):
        raise ValueError(
            f"Last populated row, column B is not a nonnegative whole count: {value!r}. "
            "If it is a formula, save the workbook in Excel so its result is cached."
        ) from None


def clean_file(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Bonds" not in workbook.sheetnames:
            raise ValueError("Sheet 'Bonds' was not found.")
        sheet = workbook["Bonds"]
        # Ignore potentially inflated worksheet dimensions from export software.
        sheet.reset_dimensions()
        rows = [(number, tuple(values)) for number, values in
                enumerate(sheet.iter_rows(values_only=True), 1)
                if number >= 4 and any(populated(v) for v in values)]
    finally:
        workbook.close()
    if len(rows) < 2 or rows[0][0] != 4:
        raise ValueError("Expected headers in row 4 and a final total row below them.")
    total_row, total_values = rows[-1]
    if len(total_values) < 2:
        raise ValueError("The final populated row has no value in column B.")
    if len(total_values) > 2 and keep_maturity(total_values[2]):
        raise ValueError("The final populated row looks like a bond, not a total row.")
    expected = parse_total(total_values[1])
    width = max(10, max(len(values) for _, values in rows))
    raw_headers = rows[0][1]
    headers, used = [], set(META + ["exclusion_reason"])
    for i in range(width):
        value = raw_headers[i] if i < len(raw_headers) else None
        base = str(value).strip() if populated(value) else f"unnamed_column_{i + 1}"
        name, suffix = base, 2
        while name in used:
            name = f"{base}__{suffix}"
            suffix += 1
        headers.append(name)
        used.add(name)
    kept, excluded = [], []
    for number, values in rows[1:-1]:
        values = list(values) + [None] * (width - len(values))
        record = dict(zip(headers, values))
        record.update(source_file=os.path.basename(path), source_excel_row=number,
                      export_isin=str(values[9]).strip().upper()
                      if populated(values[9]) else "")
        if keep_maturity(values[2]):
            kept.append(record)
        else:
            record["exclusion_reason"] = "Column C is neither a valid date nor Perpetual"
            excluded.append(record)
    matches = len(kept) == expected
    for record in kept + excluded:
        record["count_matches_total"] = matches
    ids = Counter(record["export_isin"] for record in kept if record["export_isin"])
    summary = dict(
        source_file=os.path.basename(path), status="matched" if matches else "count_mismatch",
        expected_total=expected, retained_rows=len(kept), difference=len(kept) - expected,
        excluded_nonempty_rows=len(excluded), total_excel_row=total_row,
        missing_isin_rows=sum(not record["export_isin"] for record in kept),
        duplicate_isin_rows=sum(n for n in ids.values() if n > 1),
        duplicate_isin_excess_rows=sum(n - 1 for n in ids.values()), error="",
    )
    return (pd.DataFrame(kept, columns=headers + META),
            pd.DataFrame(excluded, columns=headers + META + ["exclusion_reason"]), summary)


def save_csv(frame, path):
    temporary = path + ".tmp"
    frame.to_csv(temporary, index=False, encoding="utf-8-sig")
    os.replace(temporary, path)


def output_name(filename):
    stem = os.path.splitext(filename)[0]
    prefix = re.split(r"[0-9]", stem, maxsplit=1)[0].rstrip()
    if prefix.endswith("_"):
        prefix = prefix[:-1]
    if not prefix or prefix in (".", ".."):
        raise ValueError(f"Cannot derive a company output name from {filename!r}.")
    return prefix + ".csv"


def main():
    source = os.path.join(directory, input_folder)
    destination = os.path.join(directory, output_folder)
    if os.path.realpath(source) == os.path.realpath(destination):
        raise ValueError("Input and output folders must differ.")
    if not os.path.isdir(source):
        os.makedirs(source, exist_ok=True)
        raise RuntimeError(f"Created {source}. Place company .xlsx files there, then rerun.")
    files = sorted(name for name in os.listdir(source)
                   if name.lower().endswith((".xlsx", ".xlsm")) and not name.startswith("~$"))
    if not files:
        raise RuntimeError(f"No .xlsx or .xlsm files found in {source}.")
    names = {name: output_name(name) for name in files}
    seen = {"validation_summary.csv", "rows_excluded.csv"}
    for name, output in names.items():
        if output.casefold() in seen:
            raise ValueError(f"Output filename collision: {name!r} maps to {output!r}. "
                             "Rename the input files so each company prefix is unique.")
        seen.add(output.casefold())
    os.makedirs(destination, exist_ok=True)
    excluded_frames, summaries = [], []
    retained_total = 0
    for name in files:
        try:
            kept, excluded, summary = clean_file(os.path.join(source, name))
        except Exception as exc:
            summary = dict(source_file=name, status="error", error=str(exc))
            print(f"{name}: ERROR — {exc}")
        else:
            save_csv(kept, os.path.join(destination, names[name]))
            summary["output_file"] = names[name]
            retained_total += len(kept)
            excluded_frames.append(excluded)
            print(f"{name}: expected {summary['expected_total']:,}; "
                  f"retained {summary['retained_rows']:,}; "
                  f"difference {summary['difference']:+,}")
        summaries.append(summary)
    excluded = pd.concat(excluded_frames, ignore_index=True) if excluded_frames else pd.DataFrame(columns=META + ["exclusion_reason"])
    save_csv(excluded, os.path.join(destination, "rows_excluded.csv"))
    save_csv(pd.DataFrame(summaries), os.path.join(destination, "validation_summary.csv"))
    print(f"Outputs saved to {destination}")
    failures = sum(item["status"] != "matched" for item in summaries)
    if failures:
        raise RuntimeError(
            f"{failures} file(s) failed validation. Inspect validation_summary.csv and "
            "rows_excluded.csv before using the company files."
        )
    print(f"All {len(files)} file totals match. Retained {retained_total:,} rows; no deduplication.")


if __name__ == "__main__":
    main()
