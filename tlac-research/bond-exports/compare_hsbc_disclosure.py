"""Download HSBC's dated reference list and compare ISINs with the enriched CSV.

Run in Spyder. Requires pandas and openpyxl; no Workspace session is needed.
"""
# ---- EDIT THIS --------------------------------------------------------------
input_filename = "HSBC_Holdings_tlac.csv"
disclosure_date = "2025-12-31"  # Also supported: "2026-06-30".
isin_column = "ISIN"
issuer_column = "Issuer"
flag_column = "tlac_eligible"
# -----------------------------------------------------------------------------
try:
    from config import DATA_DIR
except ImportError:
    raise RuntimeError(
        "config.py not found. Copy config.example.py to config.py and set DATA_DIR."
    ) from None

directory = DATA_DIR.replace("\\", "/")

import os
import re
import urllib.request
from datetime import datetime, timezone

import openpyxl
import pandas as pd

BASE = "https://www.hsbc.com/-/files/hsbc/investors/hsbc-results/"
SOURCES = {
    "2025-12-31": BASE + "2025/annual/pdfs/hsbc-holdings-plc/260225-capital-and-other-tlac-eligible-instruments-main-features-31-december-2025-excel.xlsx",
    "2026-06-30": BASE + "2026/interim/pdfs/hsbc-holdings-plc/260810-capital-and-other-tlac-eligible-instruments-main-features-30-june-2026-excel.xlsx",
}
SECTIONS = {
    "CCA-CET1": "Ordinary shares",
    "CCA-AT1": "AT1: counts toward MREL",
    "CCA-T2": "Tier 2: counts toward MREL",
    "CCA-T2 (2)": "Tier 2: does not count toward MREL",
    "CCA-EL": "Other eligible liabilities: counts toward MREL",
}
FIELDS = {
    "1": "hsbc_issuer", "4": "hsbc_regulatory_treatment",
    "7": "hsbc_instrument_type", "8": "hsbc_recognised_amount_text",
    "9": "hsbc_nominal_amount_text", "11": "hsbc_issue_date_text",
    "13": "hsbc_maturity_date_text",
}


def clean(value):
    return "" if value is None else str(value).strip()


def valid_isin(value):
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", value):
        return False
    digits = "".join(str(int(c, 36)) for c in value)
    return sum(int(c) if i % 2 == 0 else sum(map(int, str(2 * int(c))))
               for i, c in enumerate(reversed(digits))) % 10 == 0


def extract_reference(path, date, url):
    workbook = openpyxl.load_workbook(path, data_only=True)
    rows = []
    if not set(SECTIONS).issubset(workbook.sheetnames):
        raise ValueError("HSBC workbook layout changed: expected category sheets missing.")
    for sheet_name, category in SECTIONS.items():
        sheet = workbook[sheet_name]
        numbered = {clean(sheet.cell(r, 1).value): r for r in range(1, sheet.max_row + 1)
                    if clean(sheet.cell(r, 1).value) in set(FIELDS) | {"2"}}
        if not (set(FIELDS) | {"2"}).issubset(numbered):
            raise ValueError(f"Missing numbered fields in {sheet_name}.")
        if "Unique identifier" not in clean(sheet.cell(numbered["2"], 2).value):
            raise ValueError(f"Unexpected identifier row in {sheet_name}.")
        for col in range(4, sheet.max_column + 1):
            raw = clean(sheet.cell(numbered["2"], col).value)
            if not raw:
                continue
            # Some published identifiers carry a superscript footnote marker.
            isin = re.sub(r"\s*[⁰¹²³⁴⁵⁶⁷⁸⁹]+$", "", raw.upper()).strip()
            if not valid_isin(isin):
                raise ValueError(f"Unrecognised identifier at {sheet_name}, column {col}: {raw}")
            row = {"comparison_isin": isin, "hsbc_category": category,
                   "hsbc_section_heading": clean(sheet.cell(2, 1).value),
                   "hsbc_disclosure_date": date, "hsbc_source_url": url,
                   "hsbc_source_sheet": sheet_name,
                   "hsbc_source_cell": sheet.cell(numbered["2"], col).coordinate}
            row.update({name: clean(sheet.cell(numbered[key], col).value)
                        for key, name in FIELDS.items()})
            rows.append(row)
    workbook.close()
    result = pd.DataFrame(rows)
    if result.empty or result.comparison_isin.duplicated().any():
        raise ValueError("Empty reference or repeated ISINs; inspect source before matching.")
    return result


def column(frame, requested):
    matches = [c for c in frame if c.strip().casefold() == requested.strip().casefold()]
    if len(matches) != 1:
        raise ValueError(f"Cannot find unique column {requested!r}. Available: {list(frame)}")
    return matches[0]


def compare(frame, reference):
    isin_name, issuer_name, flag_name = [column(frame, c) for c in
                                        [isin_column, issuer_column, flag_column]]
    data = frame.copy()
    if "comparison_isin" in data:
        raise ValueError("Use the original enriched CSV, not a previous comparison output.")
    data["comparison_isin"] = data[isin_name].str.strip().str.upper()
    flags = data[flag_name].str.strip().str.upper()
    if not flags.isin(["Y", "N", ""]).all():
        raise ValueError(f"Unexpected flag values: {flags.unique().tolist()}")
    data["comparison_flag"] = flags.replace("", "null")
    # Preserve failed requests separately from actual null classifications.
    if "tlac_lookup_status" in data:
        failed = data.tlac_lookup_status.isin(["request_error", "no_response", "missing_isin"])
        data.loc[failed & flags.eq(""), "comparison_flag"] = data.loc[failed & flags.eq(""), "tlac_lookup_status"]
    data["comparison_issuer"] = data[issuer_name].str.strip().replace("", "[missing issuer]")
    grouped = []
    for isin, part in data.loc[data.comparison_isin.ne("")].groupby("comparison_isin", sort=False):
        statuses = sorted(set(part.comparison_flag))
        grouped.append({"comparison_isin": isin, "lseg_input_rows": len(part),
                        "lseg_flags_seen": " | ".join(statuses),
                        "comparison_result": statuses[0] if len(statuses) == 1 else "conflicting_duplicate_flags",
                        "lseg_issuers": " | ".join(sorted(set(part.comparison_issuer)))})
    lookup = pd.DataFrame(grouped, columns=["comparison_isin", "lseg_input_rows", "lseg_flags_seen",
                                          "comparison_result", "lseg_issuers"])
    matched = reference.merge(lookup, on="comparison_isin", how="left", validate="one_to_one")
    matched["comparison_result"] = matched.comparison_result.fillna("absent_from_csv")
    matched["lseg_input_rows"] = pd.to_numeric(matched.lseg_input_rows).fillna(0).astype(int)
    annotated = data.merge(reference, on="comparison_isin", how="left", validate="many_to_one", sort=False)
    annotated["in_hsbc_disclosure"] = annotated.hsbc_category.notna()
    annotated["hsbc_category"] = annotated.hsbc_category.fillna("Not in selected disclosure")
    assert len(annotated) == len(frame)
    y_unlisted = annotated.loc[annotated.comparison_flag.eq("Y") & ~annotated.in_hsbc_disclosure]
    return matched, annotated, y_unlisted


def main():
    if disclosure_date not in SOURCES:
        raise ValueError(f"Supported disclosure dates: {list(SOURCES)}")
    destination = os.path.join(directory, "hsbc_disclosure_comparison", disclosure_date)
    os.makedirs(destination, exist_ok=True)
    url = SOURCES[disclosure_date]
    source = os.path.join(destination, "hsbc_main_features.xlsx")
    if not os.path.exists(source):
        request = urllib.request.Request(url, headers={"User-Agent": "HSBC disclosure research"})
        with urllib.request.urlopen(request, timeout=60) as response, open(source + ".tmp", "wb") as out:
            out.write(response.read())
        os.replace(source + ".tmp", source)
    reference = extract_reference(source, disclosure_date, url)
    reference.to_csv(os.path.join(destination, "hsbc_reference_isins.csv"), index=False, encoding="utf-8-sig")
    print(f"HSBC reference at {disclosure_date}: {len(reference):,} distinct ISINs")
    print(reference.groupby("hsbc_category").size().to_string())
    input_path = os.path.join(directory, "bond_exports_tlac", input_filename)
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Reference extracted successfully. CSV to compare is missing: {input_path}")
    frame = pd.read_csv(input_path, dtype=str, keep_default_na=False)
    matched, annotated, y_unlisted = compare(frame, reference)
    by_category = matched.groupby(["hsbc_category", "comparison_result"]).size().reset_index(name="distinct_isins")
    by_issuer = annotated.groupby(["comparison_issuer", "hsbc_category", "comparison_flag"]).agg(
        input_rows=("comparison_isin", "size"),
        distinct_nonblank_isins=("comparison_isin", lambda x: x[x.ne("")].nunique()),
    ).reset_index()
    # ISIN prefix is identifier geography, not investor residence or governing law.
    matched["isin_prefix"] = matched.comparison_isin.str[:2]
    by_prefix = matched.groupby(["hsbc_category", "isin_prefix", "comparison_result"]).size().reset_index(name="distinct_isins")
    outputs = {
        "reference_comparison.csv": matched,
        "lseg_annotated.csv": annotated,
        "hsbc_absent_from_csv.csv": matched.loc[matched.comparison_result.eq("absent_from_csv")],
        "hsbc_present_without_y.csv": matched.loc[~matched.comparison_result.isin(["Y", "absent_from_csv"])],
        "lseg_y_not_in_disclosure.csv": y_unlisted,
        "summary_by_category.csv": by_category,
        "summary_by_issuer.csv": by_issuer,
        "summary_by_isin_prefix.csv": by_prefix,
    }
    for filename, output in outputs.items():
        output.to_csv(os.path.join(destination, filename), index=False, encoding="utf-8-sig")
    keys = annotated.comparison_isin
    lines = [f"Comparison run (UTC): {datetime.now(timezone.utc).isoformat()}",
             f"HSBC disclosure: {disclosure_date}", f"Source: {url}",
             f"Input: {input_path}", f"CSV rows: {len(frame):,}",
             f"Rows without ISIN: {keys.eq('').sum():,}",
             f"Distinct nonblank CSV ISINs: {keys[keys.ne('')].nunique():,}",
             f"Extra duplicate ISIN rows: {keys.ne('').sum() - keys[keys.ne('')].nunique():,}",
             f"CSV ISINs flagged Y: {annotated.loc[annotated.comparison_flag.eq('Y') & keys.ne(''), 'comparison_isin'].nunique():,}",
             f"HSBC-listed ISINs found anywhere in CSV: {matched.comparison_result.ne('absent_from_csv').sum():,} / {len(reference):,}",
             f"HSBC-listed ISINs with consistent Y: {matched.comparison_result.eq('Y').sum():,}",
             f"Y ISINs absent from disclosure: {y_unlisted.loc[y_unlisted.comparison_isin.ne(''), 'comparison_isin'].nunique():,}",
             "", "Reference categories versus LSEG results (distinct ISINs):", by_category.to_string(index=False),
             "", "Issuer patterns (CSV rows and distinct ISINs):", by_issuer.to_string(index=False),
             "", "Ordinary shares are retained separately and need not appear in a bond export.",
             "The own-funds-only Tier 2 category is NOT treated as MREL eligible.",
             "Null and failed requests are unknown, not N. Conflicting duplicates need inspection.",
             "Missing from this dated disclosure does not establish ineligibility.",
             "Reporting-date differences, maturities, calls and new issuance can explain differences.",
             "ISIN prefixes do not identify investor residence. No monetary amounts are summed.",
             f"Outputs: {destination}"]
    if "tlac_pulled_at_utc" in frame:
        timestamps = frame.tlac_pulled_at_utc[frame.tlac_pulled_at_utc.ne("")]
        if len(timestamps):
            lines.insert(4, f"LSEG retrieval timestamps: {timestamps.min()} through {timestamps.max()}")
    report = "\n".join(lines)
    print("\n" + report)
    with open(os.path.join(destination, "summary.txt"), "w", encoding="utf-8") as stream:
        stream.write(report + "\n")


if __name__ == "__main__":
    main()
